"""RD-A L2+suggested 推荐方案：规划设备清单导出（Excel，中文）。

回答「推荐方案规划了哪些设备、每台设备的详细信息」，数据全部来自成果（`rda.suggested()` / `rda.facilities_geojson("OPTCS2")`），不另编口径。
六页：
  ① 汇总        方案说明 + 按措施类型汇总（数量 / 池容 / 泵台数 / 设计流量 / 占位造价）+ 设备数量
  ② 措施清单    53 项措施的全部台账字段（位置、截流井、池容、泵、压力管、排口、占位造价、可建性、需审批、缺数据判不了）
  ③ 设备明细    逐台设备（地下调蓄池 / 重力管 / 泵站 + 压力输水 / 地面取水口），带所属措施、服务热点和关键参数
  ④ 地面取水口  17 个取水口（地面洼坑 / 路缘）与所属地面泵站
  ⑤ 叠加的既有设施  推荐方案是叠在这些设施上算的：L1 现状泵站、L2 雨前预置移动泵、客户计划泵站（含本场模拟表现）
  ⑥ 与当前方案差异  相对「按 96 mm 优化」的当前方案，增 / 删 / 改了哪些措施

造价是占位单价，用于方案比较，不是工程概算；位置已按 WA-726 规则、用地和地下管线约束筛过，落地前仍需现场核实和主管部门许可。

用法：PYTHONPATH=src python -m stormwater.rda_sugg_equipment_xlsx [输出路径]
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import rda

NAVY = PatternFill("solid", fgColor="0F2A4A")
SOFT = PatternFill("solid", fgColor="EEF2F7")
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")

STATE = "OPTCS2"
SUGG_LEVEL = "L2+suggested 推荐方案"
#: 推荐方案新增的设备类型（设施图层 kind）；其余 kind 是叠加的既有设施
NEW_KINDS = ("tank", "gpipe", "conv", "sump")
EXISTING_KINDS = ("ps_l1", "mobile", "ps_planned")


def _put(ws, row: int, values: list, *, head: bool = False, fill=None) -> None:
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.border = BORDER
        cell.alignment = WRAP
        if head:
            cell.fill, cell.font = NAVY, HEAD_FONT
        elif fill:
            cell.fill = fill


def _table(ws, row: int, head: list[str], rows: list[list], widths: list[int] | None = None) -> int:
    """从 row 行起写一张带表头的表，返回下一个空行。"""
    _put(ws, row, head, head=True)
    for i, r in enumerate(rows, 1):
        _put(ws, row + i, r, fill=SOFT if i % 2 == 0 else None)
    if widths:
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
    return row + len(rows) + 2


def _n(v, d: int = 2):
    return None if v is None else round(float(v), d)


def _dn(m) -> str:
    return "" if m is None else f"DN{round(float(m) * 1000)}"


def _detail(kind: str, m: dict | None, intake: dict | None, p: dict) -> str:
    """一台设备的关键参数，一句话；字段取自所属措施 / 取水口 / 设施图层，缺的不编。"""
    if kind == "tank" and m:
        bits = [f"池容 {m['tank_m3']:,.0f} m³" if m.get("tank_m3") else None,
                f"有效水深 {m['effective_depth_m']} m" if m.get("effective_depth_m") is not None else None,
                f"池底 {m['invert_m']} m" if m.get("invert_m") is not None else None,
                f"开挖 {m['excavation_m']} m" if m.get("excavation_m") is not None else None,
                f"截流井 {m['intercept_wells']}" if m.get("intercept_wells") else None]
    elif kind == "gpipe" and m:
        bits = [f"重力管 截流井 → 地下池", f"长 {m['pipe_len_m']:,.0f} m" if m.get("pipe_len_m") is not None else None, _dn(m.get("pipe_dia_m")) or None]
    elif kind == "conv" and m:
        par = f" × {m['pipe_parallel']} 根" if m.get("pipe_parallel") else ""
        bits = [f"{m['pumps']} 台泵 × {m['pump_q_m3s']} m³/s" if m.get("pumps") else None,
                f"设计流量 {m['design_q_m3s']} m³/s" if m.get("design_q_m3s") is not None else None,
                f"压力管 {m['pipe_len_m']:,.0f} m {_dn(m.get('pipe_dia_m'))}{par}" if m.get("pipe_len_m") else None,
                f"流速 {m['velocity_ms']} m/s" if m.get("velocity_ms") is not None else None,
                f"排口 {m['outfall']}" if m.get("outfall") else None,
                f"泵站场地 {m['station_area_m2']:,.0f} m²" if m.get("station_area_m2") else None]
    elif kind == "sump" and intake:
        bits = [intake["类别"], f"坑底 {intake['坑底高程 m']} m", f"取水 {intake['取水流量 m³/s']} m³/s", f"归地面泵站 {intake['所属地面泵站']}"]
    else:
        bits = [p.get("说明")]
    return "；".join(b for b in bits if b)


def build() -> Workbook:
    s = rda.suggested()
    ms = s["measures"]
    mid = {m["id"]: m for m in ms}
    intake = {r["取水口"]: r for r in s["intakes"]}
    kinds = rda.FACILITY_KINDS
    feats = [f["properties"] for f in rda.facilities_geojson(STATE)["features"]]
    wb = Workbook()

    # ---- ① 汇总 ------------------------------------------------------------------
    ws = wb.active
    ws.title = "汇总"
    ws.cell(row=1, column=1, value="RD-A L2+suggested 推荐方案：规划设备清单").font = Font(bold=True, size=14)
    ws.cell(row=2, column=1, value=(
        f"推荐方案共 {len(ms)} 项措施。{s['note']}\n"
        "推演场景：RD-A，61 mm / 1 h 设计暴雨，一二维耦合（SWMM + ANUGA）；推荐方案叠加在既有设施（L1 现状泵站、L2 雨前预置移动泵、客户计划泵站）之上。"
    )).alignment = WRAP
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=8)
    ws.row_dimensions[2].height = 64
    by_type: dict[str, list[dict]] = defaultdict(list)
    for m in ms:
        by_type[m["type"]].append(m)
    cost_total = sum(m["cost_usd"] or 0 for m in ms)
    basis = {r["设施类型"]: r.get("造价依据") for r in (s["by_type"]["rows"] if s.get("by_type") else [])}   # 成果里有的类型才有依据列
    for m in ms:                       # 台账每项措施自己带了造价依据：汇总里没有时取同类型第一项的
        basis.setdefault(m["type"], m.get("cost_basis"))
        basis[m["type"]] = basis[m["type"]] or m.get("cost_basis")
    rows = []
    for t, lst in sorted(by_type.items(), key=lambda kv: -sum(m["cost_usd"] or 0 for m in kv[1])):
        c = sum(m["cost_usd"] or 0 for m in lst)
        rows.append([t, len(lst), _n(sum(m["tank_m3"] or 0 for m in lst), 0) or None, sum(m["pumps"] or 0 for m in lst) or None,
                     _n(sum(m["design_q_m3s"] or 0 for m in lst)) or None, _n(c / 1e6), f"{c / cost_total:.1%}" if cost_total else None, basis.get(t)])
    rows.append(["合计", len(ms), _n(sum(m["tank_m3"] or 0 for m in ms), 0), sum(m["pumps"] or 0 for m in ms),
                 _n(sum(m["design_q_m3s"] or 0 for m in ms)), _n(cost_total / 1e6), "100%", None])
    nxt = _table(ws, 4, ["措施类型", "数量", "池容合计 m³", "泵台数合计", "设计流量合计 m³/s", "占位造价 百万美元", "占比", "造价依据"], rows, [30, 8, 14, 12, 16, 16, 8, 60])
    cnt = defaultdict(int)
    for p in feats:
        cnt[p["kind"]] += 1
    rows = [[kinds.get(k, k), cnt[k], "推荐方案新增" if k in NEW_KINDS else "叠加的既有 / 客户计划设施"] for k in (*NEW_KINDS, *EXISTING_KINDS) if cnt[k]]
    rows.append(["可建性", None, "；".join(f"{b} {n} 项" for b, n in sorted(((b, sum(1 for m in ms if m['buildability'] == b)) for b in {m['buildability'] for m in ms}), key=lambda x: -x[1]))])
    _table(ws, nxt, ["设备 / 设施类型（设施图层）", "数量", "说明"], rows)

    # ---- ② 措施清单 ---------------------------------------------------------------
    ws = wb.create_sheet("措施清单")
    led = rda.ledger()
    head = list(led[0])
    _table(ws, 1, head, [[r.get(h) for h in head] for r in led],
           [10, 22, 10, 11, 11, 24, 10, 10, 9, 10, 9, 7, 9, 9, 9, 8, 7, 7, 9, 10, 11, 11, 12, 13, 36, 11, 40, 40, 8, 12])
    ws.freeze_panes = "C2"

    # ---- ③ 设备明细 ---------------------------------------------------------------
    ws = wb.create_sheet("设备明细")
    rows = []
    for p in feats:
        if p["kind"] not in NEW_KINDS:
            continue
        lid = rda_ledger_id(p["编号"])
        m = mid.get(lid)
        rows.append([p["编号"], lid, kinds.get(p["kind"], p["类型"]), m["hotspot"] if m else (intake.get(str(p["编号"])) or {}).get("服务热点"),
                     p.get("经度"), p.get("纬度"), p.get("流量 L/s"), p.get("池容 m³"),
                     _detail(p["kind"], m, intake.get(str(p["编号"])), p), p.get("排到 经度"), p.get("排到 纬度"), m["buildability"] if m else None, p.get("说明")])
    rows.sort(key=lambda r: (str(r[1]), str(r[0])))
    _table(ws, 1, ["设备编号", "所属措施", "设备类型", "服务热点", "经度", "纬度", "流量 L/s", "池容 m³", "关键参数", "排到 经度", "排到 纬度", "可建性", "说明"], rows,
           [11, 10, 22, 10, 11, 11, 10, 10, 60, 11, 11, 11, 60])
    ws.freeze_panes = "C2"

    # ---- ④ 地面取水口 -------------------------------------------------------------
    ws = wb.create_sheet("地面取水口")
    head = list(s["intakes"][0]) if s["intakes"] else []
    extra = ["说明"] if any("说明" in r for r in s["intakes"]) else []
    _table(ws, 1, head + extra, [[r.get(h) for h in head + extra] for r in s["intakes"]], [11, 14, 10, 12, 11, 11, 12, 14, 40][: len(head + extra)])

    # ---- ⑤ 叠加的既有设施 ---------------------------------------------------------
    ws = wb.create_sheet("叠加的既有设施")
    rows = []
    for p in feats:
        if p["kind"] not in EXISTING_KINDS:
            continue
        rows.append([p["防线层级"], p["编号"], kinds.get(p["kind"], p["类型"]), p.get("p_hotspot"), p.get("经度"), p.get("纬度"), p.get("流量 L/s"),
                     p.get("p_size"), p.get("p_on"), p.get("p_off"), p.get("p_path"), p.get("排到 经度"), p.get("排到 纬度"),
                     p.get("本场模拟_vol_m3"), p.get("本场模拟_util_pct"), p.get("说明")])
    rows.sort(key=lambda r: (r[0], str(r[1])))
    _table(ws, 1, ["防线层级", "编号", "类型", "服务热点", "经度", "纬度", "流量 L/s", "泵规格", "开泵水位 m", "停泵水位 m", "吸水到排出管长 m", "排到 经度", "排到 纬度",
                   "本场累计处理水量 m³", "利用率 %", "说明"], rows, [20, 11, 16, 10, 11, 11, 10, 8, 10, 10, 12, 11, 11, 14, 9, 70])
    ws.freeze_panes = "C2"

    # ---- ⑥ 与当前方案差异 ---------------------------------------------------------
    ws = wb.create_sheet("与当前方案差异")
    d = s.get("diff") or []
    head = list(d[0]) if d else ["变化", "编号", "类型", "服务热点", "造价 USD"]
    _table(ws, 1, head, [[r.get(h) for h in head] for r in d], [8, 10, 24, 12, 14])
    ws.cell(row=len(d) + 3, column=1, value="「当前方案」= 按 96 mm 优化的方案；表里列的是推荐方案相对它增 / 删 / 改的措施。")
    return wb


def rda_ledger_id(fid) -> str:
    """设施编号 → 措施编号（重力管 TG0496_g → TG0496，地面取水口 SP01_0 → SP01），与前端 `ledgerId` 同规则。"""
    return re.sub(r"^(SP\d+)_\d+$", r"\1", re.sub(r"_g$", "", str(fid)))


def main() -> None:
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "RD-A_L2suggested_规划设备清单.xlsx")
    wb = build()
    wb.save(out)
    print("已写出", out.resolve(), "；页：", " / ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
